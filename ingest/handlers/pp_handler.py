from __future__ import annotations

import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import numpy as np
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from config.manager import get_config
from core.constants import DOC_TYPE_SIGNALS, IMAGE_DIR, DocType
from core.types import Chunk, _token_count

SECTION_HEADER_LAYOUTS = {
    "section header",
    "section",
    "divider",
    "chapter",
}

SKIP_LAYOUTS = {"blank"}

SLIDE_TOKEN_LIMIT = 300


@dataclass
class PPTX_META:
    doc_id: str
    title: str
    doc_type: str = "GENERAL"
    slide_count: int = 0


@dataclass
class SlideElement:
    doc_id: str
    element_id: int
    slide_number: int
    type: str
    style: str = ""
    caption: str = ""
    text: str = ""
    image_path: str = ""
    x: int = 0
    y: int = 0
    height: int = 0
    width: int = 0


@dataclass
class SlideData:
    doc_id: str
    slide_number: int
    title: str
    section_name: str
    layout_name: str
    prev_slide: int = 0
    next_slide: int = 0
    elements: List[SlideElement] = field(default_factory=list)


def is_title(shape) -> bool:
    if not shape.is_placeholder:
        return False
    else:
        pp_type = shape.placeholder_format.type
        return pp_type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)


def extract_text(shape) -> str:
    lines = []
    if not shape.has_text_frame:
        return " "

    for para in shape.text_frame.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)

    return "\n".join(lines)


def shape_center(x: int, y: int, width: int, height: int) -> Tuple[float, float]:

    return (x + width / 2, y + height / 2)


def table_to_string(table) -> str:
    rows = []

    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        cells_ = "|".join(cells)
        rows.append(cells_)
    table_str = "\n".join(rows)
    return table_str


def find_nearest_text(
    image_el: "SlideElement",
    body_elements: List["SlideElement"],
    max_tokens: int = 30,
) -> Optional["SlideElement"]:

    if not body_elements:
        return None
    img_cx, img_cy = shape_center(
        image_el.x, image_el.y, image_el.width, image_el.height
    )

    best_el = None
    best_dist = np.inf

    for el in body_elements:
        if _token_count(el.text) > max_tokens:
            continue
        el_cx, el_cy = shape_center(el.x, el.y, el.width, el.height)
        distance = np.linalg.norm([img_cx - el_cx, img_cy - el_cy])
        if distance < best_dist:
            best_dist = distance
            best_el = el

    return best_el


def iter_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)
        else:
            yield shape


def classify_presentation(meta: PPTX_META, slides: List[SlideData]) -> DocType:

    high_weight = meta.title.lower()

    body_weights = " ".join(s.title for s in slides[:7]).lower()

    body_weights += (
        " "
        + " ".join(
            el.text
            for slide in slides[:7]
            for el in slide.elements
            if el.type == "body"
        ).lower()
    )

    scores: Dict[DocType, int] = {dt: 0 for dt in DocType}

    for doc_type, keywords in DOC_TYPE_SIGNALS.items():
        for kw in keywords:
            if kw in high_weight:
                scores[doc_type] += 2
            if kw in body_weights:
                scores[doc_type] += 1

    best = max(scores, key=lambda dt: scores[dt])
    best_score = scores[best]
    return best if best_score > 0 else DocType.GENERAL


def extract_slides(pptx_path: str) -> Tuple[List[SlideData], PPTX_META]:

    path = Path(pptx_path)
    doc_id = path.name

    pp_file = Presentation(pptx_path)

    try:
        title = pp_file.core_properties.title or " "
    except Exception:
        title = " "

    pp_meta = PPTX_META(doc_id=doc_id, title=title, slide_count=len(pp_file.slides))

    slides: List[SlideData] = []
    current_section = ""
    element_counter = 0
    image_counter = 0

    for slide_index, slide in enumerate(pp_file.slides):
        slide_number = slide_index + 1
        layout_names = slide.slide_layout.name.lower().strip()

        if any(s in layout_names for s in SECTION_HEADER_LAYOUTS):
            for shape in iter_shapes(slide.shapes):
                if is_title(shape):
                    section = extract_text(shape).strip()
                    if section:
                        current_section = section
                        break

        skip = any(s in layout_names for s in SKIP_LAYOUTS)

        slide_elements: List[SlideElement] = []
        title: str = ""
        if not skip:
            for shape in iter_shapes(slide.shapes):
                if is_title(shape):
                    title = extract_text(shape).strip()
                    continue

                if shape.has_table:
                    element_counter += 1
                    table_str = table_to_string(shape.table)

                    tb_element = SlideElement(
                        doc_id=doc_id,
                        element_id=element_counter,
                        text=table_str,
                        slide_number=slide_number,
                        type="table",
                    )
                    slide_elements.append(tb_element)
                    continue

                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    element_counter += 1
                    image_counter += 1

                    image_bytes = shape.image.blob
                    image_ext = shape.image.ext

                    file_name = f"{doc_id}_slide{slide_number}_image {image_counter}.{image_ext}"
                    image_path = os.path.join(IMAGE_DIR, file_name)

                    with open(image_path, "wb") as f:
                        f.write(image_bytes)

                    im_element = SlideElement(
                        doc_id=doc_id,
                        element_id=element_counter,
                        image_path=image_path,
                        slide_number=slide_number,
                        type="image",
                        x=shape.left or 0,
                        y=shape.top or 0,
                        width=shape.width or 0,
                        height=shape.height or 0,
                    )
                    slide_elements.append(im_element)

                if shape.has_text_frame:
                    element_counter += 1
                    text = extract_text(shape).strip()

                    body_element = SlideElement(
                        type="body",
                        text=text,
                        slide_number=slide_number,
                        element_id=element_counter,
                        doc_id=doc_id,
                        x=shape.left or 0,
                        y=shape.top or 0,
                        width=shape.width or 0,
                        height=shape.height or 0,
                    )
                    slide_elements.append(body_element)

        if slide.has_notes_slide:
            note_text = slide.notes_slide.notes_text_frame.text.strip()
            element_counter += 1
            note_el = SlideElement(
                type="notes",
                text=note_text,
                slide_number=slide_number,
                element_id=element_counter,
                doc_id=doc_id,
            )

            slide_elements.append(note_el)

        body_els = [e for e in slide_elements if e.type == "body"]
        claimed = set()

        for el in slide_elements:
            if el.type != "image":
                continue
            candidate = find_nearest_text(
                el,
                [b for b in body_els if b.element_id not in claimed],
            )
            if candidate:
                el.caption = candidate.text
                claimed.add(candidate.element_id)

        slide_elements = [e for e in slide_elements if e.element_id not in claimed]
        slide_el = SlideData(
            title=title,
            doc_id=doc_id,
            slide_number=slide_number,
            elements=slide_elements,
            section_name=current_section,
            layout_name=layout_names,
        )
        slides.append(slide_el)

    for index, slide in enumerate(slides):
        slide.prev_slide = slides[index - 1].slide_number if index > 0 else 0
        slide.next_slide = (
            slides[index + 1].slide_number if index < len(slides) - 1 else 0
        )

    return slides, pp_meta


def slide_section_path(slide: SlideData) -> List[str]:
    """Build section_path equivalent: [section_name, slide_title] (non-empty only)."""
    return [p for p in [slide.section_name, slide.title] if p]


def pick_size(doc_type: DocType) -> int:

    return {
        DocType.TECHNICAL: 300,
        DocType.LEGAL: 300,
        DocType.ACADEMIC: 400,
        DocType.REPORT: 350,
    }.get(doc_type, SLIDE_TOKEN_LIMIT)


def chunk(
    slides: List[SlideData], doc_meta: PPTX_META, doctype: DocType
) -> Tuple[List[Chunk], Dict[tuple, List[Chunk]], Dict[tuple, List[Chunk]]]:
    chunks: List[Chunk] = []
    image_to_chunks: Dict[tuple, List[Chunk]] = {}
    table_to_chunks: Dict[tuple, List[Chunk]] = {}
    chunk_index = 0
    doc_id = doc_meta.doc_id
    size_limit = pick_size(doctype)

    for slide in slides:
        section_path = slide_section_path(slide)
        slide_prefix = f"[SLIDE {slide.slide_number}]"
        if slide.title:
            slide_prefix += f" {slide.title}"
        if slide.section_name:
            slide_prefix += f" | {slide.section_name}"

        body_elements = [e for e in slide.elements if e.type == "body"]
        notes_elements = [e for e in slide.elements if e.type == "notes"]
        table_elements = [e for e in slide.elements if e.type == "table"]
        image_elements = [e for e in slide.elements if e.type == "image"]

        body_text = "\n".join(l.text for l in body_elements)
        full_text = (
            f"{section_path}|\n {body_text}".strip() if body_text else slide_prefix
        )

        if _token_count(full_text) <= size_limit:
            chunk = Chunk(
                text=full_text,
                doc_id=doc_id,
                chunk_index=chunk_index,
                chunk_type="text",
                style="slide_body",
                section_path=section_path,
                start_element_id=body_elements[0].element_id if body_elements else 0,
                end_element_id=body_elements[-1].element_id if body_elements else 0,
                token_count=_token_count(full_text),
            )
            chunks.append(chunk)
            chunk_index += 1
        else:
            splitter = RecursiveCharacterTextSplitter(
                chunk_size=get_config().chunking.max_tokens,
                chunk_overlap=get_config().chunking.overlap_tokens,
                separators=["\n", ". ", " ", ""],
            )

            box_texts = [f"{slide_prefix}\n{body_elements[0].text}"] + [
                e.text for e in body_elements[1:]
            ]
            joined = "\n\n".join(box_texts)
            sub_chunks = splitter.split_text(joined)

            for sp in sub_chunks:
                chunk = Chunk(
                    text=sp,
                    doc_id=doc_id,
                    chunk_index=chunk_index,
                    chunk_type="text",
                    style="slide_body",
                    section_path=section_path,
                    start_element_id=body_elements[0].element_id,
                    end_element_id=body_elements[-1].element_id,
                    token_count=_token_count(sp),
                )
                chunks.append(chunk)
                chunk_index += 1

        for el in notes_elements:
            notes_text = f"{slide_prefix} [NOTES]\n{el.text}"
            chunks.append(
                Chunk(
                    text=notes_text,
                    doc_id=doc_id,
                    chunk_index=chunk_index,
                    chunk_type="notes",
                    style="slide_notes",
                    section_path=section_path,
                    start_element_id=el.element_id,
                    end_element_id=el.element_id,
                    token_count=_token_count(notes_text),
                )
            )
            chunk_index += 1

        for el in table_elements:
            rows = el.text.split("\n")
            context = slide_prefix

            if len(rows) <= 15:
                full_text = f"{context}\n[TABLE_{el.element_id}]\n{el.text}".strip()
                c = Chunk(
                    text=full_text,
                    doc_id=doc_id,
                    chunk_index=chunk_index,
                    chunk_type="table",
                    style="slide_table",
                    section_path=section_path,
                    start_element_id=el.element_id,
                    end_element_id=el.element_id,
                    table_refs={el.element_id: el.text},
                    token_count=_token_count(full_text),
                )
                chunks.append(c)
                table_to_chunks.setdefault((doc_id, el.element_id), []).append(c)
                chunk_index += 1

            else:
                header_row = rows[0]
                data_rows = rows[1:]
                window, step = 10, 8

                for w in range(0, len(data_rows), step):
                    batch = [header_row] + data_rows[w : w + window]
                    prefix = context if w == 0 else f"[TABLE_{el.element_id} continued]"
                    full_text = f"{prefix}\n" + "\n".join(batch)

                    c = Chunk(
                        text=full_text.strip(),
                        doc_id=doc_id,
                        chunk_index=chunk_index,
                        chunk_type="table",
                        style="slide_table",
                        section_path=section_path,
                        start_element_id=el.element_id,
                        end_element_id=el.element_id,
                        table_refs={el.element_id: el.text},
                        token_count=_token_count(full_text),
                    )
                    chunks.append(c)
                    table_to_chunks.setdefault((doc_id, el.element_id), []).append(c)
                    chunk_index += 1

        for el in image_elements:
            context_text = "\n".join(body_text[:3])

            if len(context_text.split()) > 100:
                context_text = " ".join(context_text.split()[:100])

            caption_text = f"Caption: {el.caption}" if el.caption else ""

            full_text = "\n".join(
                filter(
                    None,
                    [
                        slide_prefix,
                        f"[IMAGE_{el.element_id}]",
                        caption_text,
                        context_text,
                    ],
                )
            ).strip()

            c = Chunk(
                text=full_text,
                doc_id=doc_id,
                chunk_index=chunk_index,
                chunk_type="image_context",
                style="slide_image",
                section_path=section_path,
                start_element_id=el.element_id,
                end_element_id=el.element_id,
                image_refs={el.element_id: el.image_path},
                token_count=_token_count(full_text),
            )
            chunks.append(c)
            image_to_chunks.setdefault((doc_id, el.element_id), []).append(c)
            chunk_index += 1

    return chunks, image_to_chunks, table_to_chunks


def process_pptx(pptx_path: str):

    slides, meta = extract_slides(pptx_path)
    doc_type = classify_presentation(meta, slides)
    meta.doc_type = doc_type.value

    chunks, image_to_chunks, table_to_chunks = chunk(slides, meta, doc_type)
    return chunks, meta, image_to_chunks, table_to_chunks


if __name__ == "__main__":
    path = r"F:\university\pajoohesh\power\Farham (4).pptx"

    chunks, meta, image_to_chunks, table_to_chunks = process_pptx(path)

    print(f" Presentation : {meta.doc_id}")
    print(f"   Doc type     : {meta.doc_type}")
    print(f"   Slides       : {meta.slide_count}")
    print(f"   Chunks       : {len(chunks)}")
    print(f"   img→chunk    : { {k: len(v) for k, v in image_to_chunks.items()} }")
    print(f"   tbl→chunk    : { {k: len(v) for k, v in table_to_chunks.items()} }\n")

    for c in chunks:
        print(
            f"[{c.chunk_index:03d}] type={c.chunk_type:<14} "
            f"tokens={c.token_count:<4} "
            f"path={c.section_path}"
        )
        print(f"       preview : {c.text.strip()}")
        if c.image_refs:
            print(f"       images  : {c.image_refs}")
        if c.table_refs:
            print(f"       tables  : {list(c.table_refs.keys())}")
        print()

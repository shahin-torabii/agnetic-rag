from __future__ import annotations

import os
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from data_gathering import (
Chunk, _token_count,BaseMeta
)
@dataclass
class ExcelMeta(BaseMeta):
    num_sheet:int = 0



@dataclass
class ExcelCell:
    row:int
    col:int
    value:Any
    formula:Optional[str] = " "
    color :str=""
    style :str= ""


# @dataclass
# class MergedCell:
#     range_str: str  # e.g. "B1:C2"
#     start_row: int
#     start_col: int
#     end_row: int
#     end_col: int
#     value: Any
#     formula: Optional[str] = None


@dataclass
class Sheet:
    doc_id: str
    element_id: int
    sheet_name: str
    sheet_number: int
    headers: List[List[ExcelCell]] = field(default_factory=list)
    col_paths: Dict[int, str] = field(default_factory=dict)
    data_rows: List[List[ExcelCell]] = field(default_factory=list)
    rows_as_dicts: List[Dict[str, Any]] = field(default_factory=list)
    prev_sheet: int = 0
    next_sheet: int = 0


def safe_color(cell) -> str:
    try:
        rgb = cell.fill.fgColor.rgb
        return rgb if rgb != "00000000" else ""
    except Exception:
        return ""


def build_merge_lookup(ws: Worksheet) -> Tuple[Dict, set]:

    merge_value: Dict[Tuple[int, int], Any] = {}
    covered: set = set()

    for rng in ws.merged_cells.ranges:
        val = ws.cell(rng.min_row, rng.min_col).value
        first = True
        for (r, c) in rng.cells:
            merge_value[(r, c)] = val
            if first:
                first = False
            else:
                covered.add((r, c))

    return merge_value, covered

def effective_value(ws: Worksheet, row: int, col: int,
                     merge_value: Dict) -> Any:
    if (row, col) in merge_value:
        return merge_value[(row, col)]
    return ws.cell(row, col).value

def is_header_row(ws: Worksheet, row_num: int,
                   max_col: int, merge_value: Dict) -> bool:

    non_empty = [
        effective_value(ws, row_num, col_num, merge_value)
        for col_num in range(1, max_col+1)
        if effective_value(ws, row_num, col_num, merge_value) is not None
    ]
    if not non_empty:
        return False
    return all(isinstance(cell_val, str) for cell_val in non_empty)


def detect_header_rows(ws: Worksheet, merge_value: Dict) -> List[int]:

    header_rows = []
    for row in range(1, ws.max_row+1):
        if is_header_row(ws, row, ws.max_column, merge_value):
            header_rows.append(row)
        else:
            break

    return header_rows

def build_col_paths(ws: Worksheet, header_rows: List[int],
                     merge_value: Dict) -> Dict[int, str]:

    col_path :Dict[int, str] = {}
    for col in range(1, ws.max_column+1):
        parts = []
        seen = set()
        for hr in header_rows:
            value = effective_value(ws, hr, col, merge_value)
            if value:
                text = str(value).strip()
                if text and not text in seen:
                    parts.append(text)
                    seen.add(text)
        parts = ">".join(parts) if parts else f"col_{col}"
        col_path[col] = parts

    return col_path

def row_to_dict(cells: List[ExcelCell],
                 col_paths: Dict[int, str]) -> Dict[str, Any]:
    row_to_dict = {
            col_paths[cell.col]: cell.value
            for cell in cells
            if cell.value is not None and cell.col in col_paths
        }

    return row_to_dict



def extract_rows(excel_path: str) -> Tuple[List[Sheet], ExcelMeta]:
    path = Path(excel_path)
    doc_id = path.name

    wb_val = load_workbook(excel_path, data_only=True)
    wb_frm = load_workbook(excel_path, data_only=False)

    try:
        title = wb_val.properties.title or doc_id
    except Exception:
        title = doc_id

    doc_meta = ExcelMeta(
        doc_id=doc_id,
        title = title,
        source_type="xlsx",
        num_sheet=len(wb_val.sheetnames)
    )
    sheets:List[Sheet] = []
    element_id = 0

    for sheet_index, sheet_name in enumerate(wb_val.sheetnames):
        ws_v = wb_val[sheet_name]
        ws_f = wb_frm[sheet_name]

        merge_value, covered = build_merge_lookup(ws_v)

        header_row_nums = detect_header_rows(ws_v, merge_value)
        data_start_row = (max(header_row_nums) + 1) if header_row_nums else 1


        col_paths = build_col_paths(ws_v, header_row_nums, merge_value)


        all_rows: Dict[int, List[ExcelCell]] = {}
        for row_num in range(1, ws_f.max_row+1):
            cells :List[ExcelCell] = []
            for col_num in range(1, ws_f.max_column +1):
                if (row_num, col_num) in covered:
                    continue
                v_cell = ws_v.cell(row_num, col_num)
                f_cell = ws_f.cell(row_num, col_num)

                formula = (
                    f_cell.value
                    if isinstance(f_cell.value, str) and f_cell.value.startswith("=")
                    else None
                )

                value = effective_value(ws_v, row_num, col_num, merge_value)

                cells.append(ExcelCell(
                    row=row_num,
                    col=col_num,
                    value=value,
                    formula=formula,
                    color=safe_color(v_cell),
                    style=v_cell.style or "",
                ))

                if cells:
                    all_rows[row_num] = cells

        header_cells = [
            all_rows[r] for r in header_row_nums if r in all_rows
        ]
        data_cells = [
            all_rows[r]
            for r in range(data_start_row, ws_v.max_row + 1)
            if r in all_rows
        ]

        rows_as_dicts = [row_to_dict(row, col_paths) for row in data_cells]
        rows_as_dicts = [d for d in rows_as_dicts if d]  # drop empty rows

        sheet = Sheet(
            doc_id=doc_id,
            element_id=element_id,
            sheet_number=sheet_index,
            sheet_name=sheet_name,
            headers=header_cells,
            col_paths=col_paths,
            data_rows=data_cells,
            rows_as_dicts=rows_as_dicts,
        )
        sheets.append(sheet)
        element_id += 1


    for i, sheet in enumerate(sheets):
        sheet.prev_sheet = sheets[i - 1].sheet_number if i > 0 else 0
        sheet.next_sheet = sheets[i + 1].sheet_number if i < len(sheets) - 1 else 0

    return sheets, doc_meta

def build_row_text(row_dict):

    lines = []

    for key, value in row_dict.items():
        lines.append(f"{key}: {value}")

    return "\n".join(lines)



def chunk(
    sheets: List[Sheet],
    doc_meta: ExcelMeta
) -> Tuple[List[Chunk], ExcelMeta]:

    chunks = []

    chunk_index = 0

    for sheet in sheets:

        header_text = "\n".join(sheet.col_paths.values())

        rows = sheet.rows_as_dicts

        if not rows:
            continue


        if len(rows) <= 15:

            row_text = "\n\n".join(
                build_row_text(r)
                for r in rows
            )

            full_text = (
                f"Workbook: {doc_meta.title}\n"
                f"Sheet: {sheet.sheet_name}\n\n"
                f"Columns:\n{header_text}\n\n"
                f"{row_text}"
            )

            chunks.append(
                Chunk(
                    text=full_text,
                    doc_id=doc_meta.doc_id,
                    chunk_index=chunk_index,
                    chunk_type="table",
                    start_element_id=sheet.element_id,
                    end_element_id=sheet.element_id,
                    token_count=_token_count(full_text)
                )
            )

            chunk_index += 1

        else:

            window = 25
            step = 20

            for start in range(0, len(rows), step):

                end = min(start + window, len(rows))

                row_text = "\n\n".join(
                    build_row_text(r)
                    for r in rows[start:end]
                )

                full_text = (
                    f"Workbook: {doc_meta.title}\n"
                    f"Sheet: {sheet.sheet_name}\n\n"
                    f"Columns:\n{header_text}\n\n"
                    f"{row_text}"
                )

                chunks.append(
                    Chunk(
                        text=full_text,
                        doc_id=doc_meta.doc_id,
                        chunk_index=chunk_index,
                        chunk_type="table",
                        start_element_id=sheet.element_id,
                        end_element_id=sheet.element_id,
                        token_count=_token_count(full_text)
                    )
                )

                chunk_index += 1

    return chunks, doc_meta



def process_excel(excel_path: str):

    sheets, meta = extract_rows(excel_path)

    chunks, doc_meta= chunk(sheets, meta)
    return chunks, meta,


if __name__ == "__main__":
    path = r"C:\Users\shahin\Desktop\Logical Circuits - Fall 1404.xlsx"

    chunks, meta = process_excel(path)

    print(f" Excel: {meta.doc_id}")
    print(f"   Sheets       : {meta.num_sheet}")
    print(f"   Chunks       : {len(chunks)}")

    for c in chunks:
        print(
            f"[{c.chunk_index:03d}] type={c.chunk_type:<14} "
            f"tokens={c.token_count:<4} "
        )
        print(f"       preview : {c.text.strip()}")
        print("\n")